"""
FORGE Document Ingestion Pipeline (Phase 2)
Supported types: PDF, DOCX, JPG, PNG.

STRICT HONESTY MANDATE: This module never fabricates extracted text. Earlier
versions returned a hardcoded placeholder string ("[PDF text extraction not
wired in this reference build.]") for every file regardless of content, with
a fixed fake ocr_confidence and page_count=1 always -- meaning every
downstream completeness check and candidate-fact proposal ran against
fictional input. That has been replaced with real extraction for PDF (pypdf)
and DOCX (python-docx). For image files, OCR requires a system-level
Tesseract binary this module cannot verify is present in every deployment;
rather than fake a result, it returns an explicit UNAVAILABLE sentinel that
completeness.py checks for and surfaces as a hard flag, so the user is never
told "no issues detected" about a document nothing actually read.
"""
from __future__ import annotations
import hashlib
import io
import re
from dataclasses import dataclass
from typing import List, Optional, Tuple

SUPPORTED_MIME_TYPES = {
    "application/pdf": "pdf",
    "image/jpeg": "jpg",
    "image/png": "png",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}
MAX_FILE_SIZE_BYTES = 25 * 1024 * 1024

EXTRACTION_UNAVAILABLE_MARKER = "[FORGE_EXTRACTION_UNAVAILABLE]"
EXTRACTION_UNAVAILABLE_TEXT = (
    f"{EXTRACTION_UNAVAILABLE_MARKER} FORGE could not extract machine-readable text "
    f"from this file in the current environment. No completeness check, fact "
    f"proposal, or contradiction scan has run against this document's actual "
    f"content -- please review it yourself."
)

DATE_PATTERN = re.compile(
    r"\b(\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|"
    r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{1,2},?\s+\d{4})\b",
    re.IGNORECASE
)
CASE_NUMBER_PATTERN = re.compile(r"\b\d{1,2}:\d{2}-cv-\d{3,6}(?:-[A-Z]{2,5})?\b")
AMOUNT_PATTERN = re.compile(r"\$\s?[\d,]+(?:\.\d{2})?")


@dataclass
class ValidationResult:
    is_valid: bool
    reason: Optional[str] = None
    normalized_type: Optional[str] = None


def validate_upload(filename: str, mime_type: str, size_bytes: int) -> ValidationResult:
    if size_bytes <= 0:
        return ValidationResult(False, "The file appears to be empty.")
    if size_bytes > MAX_FILE_SIZE_BYTES:
        return ValidationResult(False, f"File exceeds the {MAX_FILE_SIZE_BYTES // (1024*1024)} MB limit for this release.")
    if mime_type not in SUPPORTED_MIME_TYPES:
        return ValidationResult(False, "This file type is not yet supported. Supported types: PDF, DOCX, JPG, PNG.")
    return ValidationResult(True, normalized_type=SUPPORTED_MIME_TYPES[mime_type])


def compute_sha256(file_bytes: bytes) -> str:
    return hashlib.sha256(file_bytes).hexdigest()


def _extract_pdf(file_bytes: bytes) -> Tuple[str, Optional[float], Optional[int]]:
    """Real extraction via pypdf. Returns (text, confidence, page_count).
    confidence here means 'did we get a real text layer', not OCR quality:
    1.0 if non-trivial text was found, a low score if the PDF appears to be
    scanned images with no embedded text layer (which would need OCR FORGE
    does not perform)."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return EXTRACTION_UNAVAILABLE_TEXT, None, None

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        page_count = len(reader.pages)
        pages_text = []
        for page in reader.pages:
            pages_text.append(page.extract_text() or "")
        full_text = "\n".join(pages_text).strip()
    except Exception:
        return EXTRACTION_UNAVAILABLE_TEXT, None, None

    if len(full_text) < 20:
        return (
            f"{EXTRACTION_UNAVAILABLE_MARKER} This PDF appears to contain scanned "
            f"images with no embedded text layer. FORGE could not extract "
            f"machine-readable text from it -- consider uploading a text-based "
            f"version or running OCR on it first.",
            0.1, page_count,
        )
    return full_text, 1.0, page_count


def _extract_docx(file_bytes: bytes) -> Tuple[str, Optional[float], Optional[int]]:
    """Real extraction via python-docx. DOCX files do not store a rendered
    page count (pagination depends on the viewer/printer), so page_count is
    returned as None rather than guessed -- completeness.py's page-count
    check already treats None as 'not applicable' rather than a red flag."""
    try:
        from docx import Document
    except ImportError:
        return EXTRACTION_UNAVAILABLE_TEXT, None, None

    try:
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)
        full_text = "\n".join(p for p in paragraphs if p).strip()
    except Exception:
        return EXTRACTION_UNAVAILABLE_TEXT, None, None

    if len(full_text) < 5:
        return EXTRACTION_UNAVAILABLE_TEXT, None, None
    return full_text, 1.0, None


def _extract_image(file_bytes: bytes) -> Tuple[str, Optional[float], Optional[int]]:
    """No OCR engine is bundled with this module: pytesseract requires a
    system-level Tesseract binary this module cannot verify is installed in
    every FORGE deployment. Rather than silently skip OCR and pretend
    nothing is wrong, this returns the explicit UNAVAILABLE sentinel so
    completeness.py raises a hard flag instead of reporting a clean bill of
    health on a document nothing actually read."""
    return EXTRACTION_UNAVAILABLE_TEXT, None, 1


def extract_text_stub(file_bytes: bytes, normalized_type: str) -> Tuple[str, float, int]:
    """Kept for backward-compatible import paths; now dispatches to real
    extraction instead of returning a hardcoded placeholder. Callers should
    treat a None confidence/page_count and an EXTRACTION_UNAVAILABLE_MARKER
    prefix as 'not actually reviewed' rather than a valid low score."""
    if normalized_type == "docx":
        text, conf, pages = _extract_docx(file_bytes)
    elif normalized_type == "pdf":
        text, conf, pages = _extract_pdf(file_bytes)
    else:
        text, conf, pages = _extract_image(file_bytes)
    return text, conf, pages


def propose_candidate_facts(extracted_text: str, page_number: int = 1) -> List[dict]:
    if extracted_text.startswith(EXTRACTION_UNAVAILABLE_MARKER):
        return []

    candidates = []
    for match in DATE_PATTERN.finditer(extracted_text):
        start, end = max(0, match.start() - 40), min(len(extracted_text), match.end() + 40)
        candidates.append({
            "normalized_statement": f"Document references date: {match.group(0)}",
            "fact_type": "date", "confidence": 0.6, "source_page": page_number,
            "source_quote": extracted_text[start:end].strip(),
        })
    for match in CASE_NUMBER_PATTERN.finditer(extracted_text):
        candidates.append({
            "normalized_statement": f"Document references case number: {match.group(0)}",
            "fact_type": "case_number", "confidence": 0.7, "source_page": page_number,
            "source_quote": match.group(0),
        })
    for match in AMOUNT_PATTERN.finditer(extracted_text):
        candidates.append({
            "normalized_statement": f"Document references amount: {match.group(0)}",
            "fact_type": "amount", "confidence": 0.55, "source_page": page_number,
            "source_quote": match.group(0),
        })
    return candidates


def run_pipeline(filename: str, mime_type: str, file_bytes: bytes) -> dict:
    validation = validate_upload(filename, mime_type, len(file_bytes))
    if not validation.is_valid:
        return {"status": "rejected", "reason": validation.reason}
    sha256 = compute_sha256(file_bytes)
    extracted_text, ocr_confidence, page_count = extract_text_stub(file_bytes, validation.normalized_type)
    candidate_facts = propose_candidate_facts(extracted_text, page_number=1)
    return {
        "status": "processed", "normalized_type": validation.normalized_type, "sha256": sha256,
        "extracted_text": extracted_text, "ocr_confidence": ocr_confidence,
        "page_count": page_count, "candidate_facts": candidate_facts,
    }
